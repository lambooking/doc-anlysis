"""审核报告生成器模块"""
from docx import Document
from docx.shared import RGBColor, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from pathlib import Path
from typing import Dict
from ..models.schemas import AuditReport, AuditResult
from ..utils.logger import get_logger

logger = get_logger()


class ReportGenerator:
    """审核报告生成器 - 生成 DOCX 格式报告"""
    
    # 严重程度标签映射
    SEVERITY_LABELS = {
        'critical': '关键问题',
        'high': '高优先级问题',
        'medium': '中优先级问题',
        'low': '低优先级问题'
    }
    
    # 严重程度颜色
    SEVERITY_COLORS = {
        'critical': RGBColor(255, 0, 0),
        'high': RGBColor(255, 128, 0),
        'medium': RGBColor(255, 255, 0),
        'low': RGBColor(128, 255, 0)
    }
    
    def __init__(self):
        pass
    
    def generate_report(
        self,
        audit_report: AuditReport,
        output_path: str
    ) -> str:
        """
        生成审核报告
        
        Args:
            audit_report: 审核报告数据
            output_path: 输出路径
        
        Returns:
            报告文件路径
        """
        logger.info(f"开始生成审核报告: {audit_report.document_name}")
        
        doc = Document()
        
        # 设置文档样式
        self._setup_document_styles(doc)
        
        # 1. 标题
        self._add_title(doc)
        
        # 2. 元数据表格
        self._add_metadata_section(doc, audit_report)
        
        # 3. 审核摘要
        self._add_summary_section(doc, audit_report)
        
        # 4. 评分结果
        self._add_score_section(doc, audit_report)
        
        # 5. 详细问题列表
        self._add_violations_section(doc, audit_report)
        
        # 6. 扣分汇总表（新增）
        self._add_deduction_breakdown(doc, audit_report)

        # 7. 建议
        self._add_recommendations_section(doc, audit_report)
        
        # 保存
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        
        logger.info(f"报告生成完成: {output_path}")
        
        return str(output_path)
    
    def _setup_document_styles(self, doc: Document):
        """设置文档样式"""
        # 设置默认字体
        style = doc.styles['Normal']
        style.font.name = '宋体'
        style.font.size = Pt(12)
    
    def _add_title(self, doc: Document):
        """添加标题"""
        title = doc.add_heading('文档审核报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()  # 空行
    
    def _add_metadata_section(self, doc: Document, report: AuditReport):
        """添加元数据部分"""
        doc.add_heading('基本信息', 1)
        
        table = doc.add_table(rows=5, cols=2)
        table.style = 'Light Grid Accent 1'
        
        # 填充数据
        table.rows[0].cells[0].text = '文档名称'
        table.rows[0].cells[1].text = report.document_name
        
        table.rows[1].cells[0].text = '审核场景'
        table.rows[1].cells[1].text = report.scenario_name
        
        table.rows[2].cells[0].text = '审核日期'
        table.rows[2].cells[1].text = report.audit_date.strftime('%Y年%m月%d日 %H:%M:%S')
        
        table.rows[3].cells[0].text = '审核人'
        table.rows[3].cells[1].text = report.auditor
        
        table.rows[4].cells[0].text = '处理时间'
        table.rows[4].cells[1].text = f"{report.processing_time:.2f} 秒"
        
        doc.add_paragraph()  # 空行
    
    def _add_summary_section(self, doc: Document, report: AuditReport):
        """添加审核摘要"""
        doc.add_heading('审核摘要', 1)
        
        result = report.audit_result
        
        # 统计各严重程度的问题数量
        severity_counts = self._count_by_severity(result)
        
        table = doc.add_table(rows=6, cols=2)
        table.style = 'Light List Accent 1'
        
        table.rows[0].cells[0].text = '关键问题'
        table.rows[0].cells[1].text = str(severity_counts['critical'])
        
        table.rows[1].cells[0].text = '高优先级问题'
        table.rows[1].cells[1].text = str(severity_counts['high'])
        
        table.rows[2].cells[0].text = '中优先级问题'
        table.rows[2].cells[1].text = str(severity_counts['medium'])
        
        table.rows[3].cells[0].text = '低优先级问题'
        table.rows[3].cells[1].text = str(severity_counts['low'])
        
        table.rows[4].cells[0].text = '总问题数'
        table.rows[4].cells[1].text = str(len(result.violations))
        
        table.rows[5].cells[0].text = '总扣分'
        table.rows[5].cells[1].text = f"{result.total_deductions} 分"
        
        doc.add_paragraph()  # 空行
    
    def _add_score_section(self, doc: Document, report: AuditReport):
        """添加评分结果"""
        doc.add_heading('评分结果', 1)
        
        result = report.audit_result
        
        # 得分段落
        p = doc.add_paragraph()
        p.add_run('最终得分: ').bold = True
        score_run = p.add_run(f"{result.final_score} / 100")
        score_run.font.size = Pt(18)
        score_run.font.bold = True
        
        # 根据分数着色
        if result.final_score >= 90:
            score_run.font.color.rgb = RGBColor(0, 128, 0)
        elif result.final_score >= 60:
            score_run.font.color.rgb = RGBColor(255, 165, 0)
        else:
            score_run.font.color.rgb = RGBColor(255, 0, 0)
        
        # 通过状态
        p = doc.add_paragraph()
        p.add_run('审核结果: ').bold = True
        status_run = p.add_run('通过 ✓' if result.passed else '不通过 ✗')
        status_run.font.size = Pt(16)
        status_run.font.bold = True
        status_run.font.color.rgb = RGBColor(0, 128, 0) if result.passed else RGBColor(255, 0, 0)
        
        doc.add_paragraph()  # 空行
    
    def _add_violations_section(self, doc: Document, report: AuditReport):
        """添加详细问题列表"""
        doc.add_heading('详细问题列表', 1)
        
        result = report.audit_result
        
        if not result.violations:
            doc.add_paragraph('未发现任何问题。')
            return
        
        # 按严重程度分组
        for severity in ['critical', 'high', 'medium', 'low']:
            items = [v for v in result.violations if v.severity == severity]
            
            if items:
                label = self.SEVERITY_LABELS[severity]
                doc.add_heading(f'{label} ({len(items)})', 2)
                
                for i, violation in enumerate(items, 1):
                    # 问题标题
                    p = doc.add_paragraph()
                    p.add_run(f"{i}. [{violation.rule_id}] ").bold = True
                    title_run = p.add_run(violation.finding)
                    title_run.font.color.rgb = self.SEVERITY_COLORS.get(
                        severity, 
                        RGBColor(0, 0, 0)
                    )
                    
                    # 位置信息
                    loc_p = doc.add_paragraph(style='List Bullet 2')
                    loc_p.add_run('位置: ').italic = True
                    loc_p.add_run(
                        f"第 {violation.location.page} 页，"
                        f"{violation.location.region_description}"
                    )
                    
                    # 文本片段
                    text_p = doc.add_paragraph(style='List Bullet 2')
                    text_p.add_run('文本片段: ').italic = True
                    text_p.add_run(f'"{violation.location.text_snippet}"')
                    
                    # 扣分
                    deduct_p = doc.add_paragraph(style='List Bullet 2')
                    deduct_p.add_run('扣分: ').italic = True
                    deduct_run = deduct_p.add_run(f"{violation.points_deducted} 分")
                    deduct_run.font.color.rgb = RGBColor(255, 0, 0)
                    
                    doc.add_paragraph()  # 空行
    
    def _add_recommendations_section(self, doc: Document, report: AuditReport):
        """添加建议部分"""
        doc.add_heading('改进建议', 1)
        
        result = report.audit_result
        severity_counts = self._count_by_severity(result)
        
        if not result.violations:
            doc.add_paragraph('文档符合所有审核要求，无需改进。')
            return
        
        # 根据问题严重程度提供建议
        if severity_counts['critical'] > 0:
            p = doc.add_paragraph('1. ', style='List Number')
            p.add_run('关键问题必须立即处理：').bold = True
            p.add_run(f"发现 {severity_counts['critical']} 个关键问题，这些问题可能导致严重后果，必须优先解决。")
        
        if severity_counts['high'] > 0:
            p = doc.add_paragraph('2. ', style='List Number')
            p.add_run('高优先级问题建议尽快处理：').bold = True
            p.add_run(f"发现 {severity_counts['high']} 个高优先级问题，建议在下次修订时解决。")
        
        if severity_counts['medium'] > 0:
            p = doc.add_paragraph('3. ', style='List Number')
            p.add_run('中优先级问题建议改进：').bold = True
            p.add_run(f"发现 {severity_counts['medium']} 个中优先级问题，建议逐步改进。")
        
        if severity_counts['low'] > 0:
            p = doc.add_paragraph('4. ', style='List Number')
            p.add_run('低优先级问题可择机处理：').bold = True
            p.add_run(f"发现 {severity_counts['low']} 个低优先级问题，可在有时间时进行优化。")
        
        doc.add_paragraph()
        doc.add_paragraph('请参考详细问题列表，逐项进行修改和完善。')
    
    def _count_by_severity(self, result: AuditResult) -> Dict[str, int]:
        """统计各严重程度的问题数量"""
        counts = {
            'critical': 0,
            'high': 0,
            'medium': 0,
            'low': 0
        }
        
        for violation in result.violations:
            counts[violation.severity] = counts.get(violation.severity, 0) + 1
        
        return counts

    def _add_deduction_breakdown(self, doc: Document, report: AuditReport):
        """新增：按维度/检查点扣分汇总表。
        说明：由于当前 `AuditResult` 不携带每维度扣分明细，这里根据
        `violation.rule_id` 反查规则配置进行聚合。
        """
        from ..models.rule_engine import RuleEngine

        doc.add_heading('扣分汇总', 1)

        re_engine = RuleEngine()  # 使用默认规则文件

        # 归并：维度 -> 检查点 -> 扣分
        dim_map = {}
        for v in report.audit_result.violations:
            # 在所有场景中查找 rule_id
            hit = None
            for sid in re_engine.get_all_scenario_ids():
                info = re_engine.get_checkpoint_info(sid, v.rule_id)
                if info:
                    hit = info
                    break
            if hit is None:
                dim_key = ('未归类', 'others')
                cp_key = v.rule_id
                max_ded = 5
            else:
                dim = hit['dimension_id']
                dim_name = dim
                dim_key = (dim_name, dim)
                cp_key = hit['checkpoint'].checkpoint_id
                max_ded = hit['checkpoint'].max_deduction

            dim_map.setdefault(dim_key, {})
            current = dim_map[dim_key].get(cp_key, 0)
            current += max(0, v.points_deducted)
            dim_map[dim_key][cp_key] = min(current, max_ded)

        if not dim_map:
            doc.add_paragraph('无扣分。')
            return

        # 输出表格
        for (dim_name, dim_id), cp_dict in dim_map.items():
            doc.add_heading(f'维度：{dim_name}', 2)
            table = doc.add_table(rows=1, cols=3)
            hdr = table.rows[0].cells
            hdr[0].text = '检查点ID'
            hdr[1].text = '扣分'
            hdr[2].text = '说明'
            table.style = 'Light Grid Accent 1'

            total_dim_ded = 0
            for cp_id, ded in sorted(cp_dict.items()):
                row = table.add_row().cells
                row[0].text = cp_id
                row[1].text = str(ded)
                row[2].text = '累计至上限封顶'
                total_dim_ded += ded

            # 维度小结
            p = doc.add_paragraph()
            p.add_run(f'维度小结扣分：{total_dim_ded} 分').bold = True
            doc.add_paragraph()

