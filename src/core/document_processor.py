"""文档处理核心模块"""
import pymupdf
import pymupdf4llm
from docx import Document
from pathlib import Path
from typing import Optional, Dict, Any, List
import subprocess
import shutil
from ..models.schemas import (
    DocumentStructure, PageStructure, ParagraphInfo, 
    TableInfo, ImageInfo
)
from ..utils.logger import get_logger

logger = get_logger()


class DocumentProcessor:
    """文档处理器 - 支持 PDF、DOCX、DOC 格式"""
    
    def __init__(self, output_dir: str = "temp"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.image_dir = self.output_dir / "images"
        self.image_dir.mkdir(exist_ok=True)
    
    def process_document(self, file_path: str) -> DocumentStructure:
        """
        主入口 - 根据格式路由处理
        
        Args:
            file_path: 文档路径
        
        Returns:
            文档结构数据
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        suffix = file_path.suffix.lower()
        
        logger.info(f"开始处理文档: {file_path.name}, 格式: {suffix}")
        
        if suffix == '.pdf':
            return self.process_pdf(file_path)
        elif suffix == '.docx':
            return self.process_docx(file_path)
        elif suffix == '.doc':
            return self.process_doc(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {suffix}")
    
    def process_pdf(self, pdf_path: Path) -> DocumentStructure:
        """
        处理 PDF 文档
        
        Args:
            pdf_path: PDF 文件路径
        
        Returns:
            文档结构数据
        """
        logger.info(f"使用 PyMuPDF 处理 PDF: {pdf_path.name}")
        
        doc = pymupdf.open(str(pdf_path))
        
        # 提取 Markdown 格式（适合 VLM 输入）
        try:
            md_result = pymupdf4llm.to_markdown(
                doc=str(pdf_path),
                page_chunks=True,
                write_images=True,
                image_path=str(self.image_dir),
                embed_images=False
            )
            
            # 处理返回结果：如果是列表则合并为字符串
            if isinstance(md_result, list):
                md_text = "\n\n".join(
                    page.get('text', '') if isinstance(page, dict) else str(page)
                    for page in md_result
                )
            else:
                md_text = str(md_result)
                
        except Exception as e:
            logger.warning(f"Markdown 提取失败，使用备用方法: {e}")
            md_text = self._extract_text_fallback(doc)
        
        # 提取结构化元数据
        structure = []
        for page_num, page in enumerate(doc):
            page_info = self._extract_page_structure(page, page_num + 1)
            structure.append(page_info)
        
        page_count = len(doc)
        doc.close()
        
        logger.info(f"PDF 处理完成: {page_count} 页")
        
        return DocumentStructure(
            markdown=md_text,
            structure=structure,
            format="pdf",
            page_count=page_count
        )
    
    def process_docx(self, docx_path: Path) -> DocumentStructure:
        """
        处理 DOCX 文档
        
        Args:
            docx_path: DOCX 文件路径
        
        Returns:
            文档结构数据
        """
        logger.info(f"处理 DOCX 文档: {docx_path.name}")
        
        # DOCX 转 PDF 后处理
        pdf_path = self._convert_docx_to_pdf(docx_path)
        
        if pdf_path and pdf_path.exists():
            result = self.process_pdf(pdf_path)
            result.format = "docx"
            return result
        else:
            # 备用方案：直接提取 DOCX 文本
            logger.warning("DOCX 转 PDF 失败，使用备用方案")
            return self._process_docx_direct(docx_path)
    
    def process_doc(self, doc_path: Path) -> DocumentStructure:
        """
        处理 DOC 文档（需要转换为 DOCX）
        
        Args:
            doc_path: DOC 文件路径
        
        Returns:
            文档结构数据
        """
        logger.info(f"处理 DOC 文档: {doc_path.name}")
        
        # DOC 转 DOCX
        docx_path = self._convert_doc_to_docx(doc_path)
        
        if docx_path and docx_path.exists():
            result = self.process_docx(docx_path)
            result.format = "doc"
            return result
        else:
            raise RuntimeError("DOC 转 DOCX 失败，请检查 LibreOffice 是否安装")
    
    def generate_page_images(self, pdf_path: str, dpi: int = 150) -> List[Dict[str, Any]]:
        """
        生成页面图片用于 VLM 输入
        
        Args:
            pdf_path: PDF 文件路径
            dpi: 图片 DPI
        
        Returns:
            页面图片信息列表
        """
        logger.info(f"生成页面图片，DPI: {dpi}")
        
        doc = pymupdf.open(pdf_path)
        page_images = []
        
        for page_num, page in enumerate(doc):
            # 生成图片
            mat = pymupdf.Matrix(dpi / 72, dpi / 72)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            
            # 保存图片
            image_filename = f"page_{page_num + 1}.png"
            image_path = self.image_dir / image_filename
            pix.save(str(image_path))
            
            page_images.append({
                'page': page_num + 1,
                'image_path': str(image_path),
                'width': pix.width,
                'height': pix.height
            })
            
            pix = None
        
        doc.close()
        
        logger.info(f"生成 {len(page_images)} 个页面图片")
        return page_images
    
    def _extract_page_structure(self, page, page_num: int) -> PageStructure:
        """提取单页结构信息"""
        page_info = PageStructure(
            page=page_num,
            paragraphs=[],
            tables=[],
            images=[]
        )
        
        # 提取文本块及位置（同时保留前若干段落内容用于签字页/日期粗定位）
        try:
            text_dict = page.get_text("dict")
            para_idx = 0
            
            for block in text_dict.get("blocks", []):
                if block.get("type") == 0:  # 文本块
                    text = " ".join(
                        s.get("text", "") for line in block.get("lines", [])
                        for s in line.get("spans", [])
                    )
                    if text.strip():
                        page_info.paragraphs.append(ParagraphInfo(
                            index=para_idx,
                            bbox=block.get("bbox"),
                            text=text.strip()
                        ))
                        para_idx += 1
        except Exception as e:
            logger.warning(f"页面 {page_num} 文本提取失败: {e}")
        
        # 提取表格
        try:
            tables = page.find_tables()
            for idx, table in enumerate(tables):
                # 提取表格数据并处理 None 值
                table_data = []
                if hasattr(table, 'extract'):
                    raw_data = table.extract()
                    if raw_data:
                        # 将 None 转换为空字符串
                        table_data = [
                            [str(cell) if cell is not None else '' for cell in row]
                            for row in raw_data
                        ]
                
                page_info.tables.append(TableInfo(
                    index=idx,
                    bbox=list(table.bbox) if hasattr(table.bbox, '__iter__') else None,
                    data=table_data
                ))
        except Exception as e:
            logger.warning(f"页面 {page_num} 表格提取失败: {e}")
        
        return page_info
    
    def _extract_text_fallback(self, doc) -> str:
        """备用文本提取方法"""
        text_parts = []
        for page_num, page in enumerate(doc):
            text_parts.append(f"\n\n--- 第 {page_num + 1} 页 ---\n\n")
            text_parts.append(page.get_text())
        return "".join(text_parts)
    
    def _convert_docx_to_pdf(self, docx_path: Path) -> Optional[Path]:
        """
        将 DOCX 转换为 PDF
        
        使用 LibreOffice 进行转换
        """
        output_pdf = self.output_dir / f"{docx_path.stem}.pdf"
        
        # 尝试使用 LibreOffice
        if shutil.which("soffice") or shutil.which("libreoffice"):
            try:
                cmd = [
                    "soffice" if shutil.which("soffice") else "libreoffice",
                    "--headless",
                    "--convert-to", "pdf",
                    "--outdir", str(self.output_dir),
                    str(docx_path)
                ]
                
                result = subprocess.run(
                    cmd,
                    timeout=30,
                    capture_output=True,
                    text=True
                )
                
                if result.returncode == 0 and output_pdf.exists():
                    logger.info(f"DOCX 转 PDF 成功: {output_pdf.name}")
                    return output_pdf
                else:
                    logger.warning(f"LibreOffice 转换失败: {result.stderr}")
            except Exception as e:
                logger.warning(f"LibreOffice 转换异常: {e}")
        
        # 尝试使用 python-docx2pdf (仅 Windows/macOS)
        try:
            from docx2pdf import convert
            convert(str(docx_path), str(output_pdf))
            if output_pdf.exists():
                logger.info(f"使用 docx2pdf 转换成功")
                return output_pdf
        except Exception as e:
            logger.warning(f"docx2pdf 转换失败: {e}")
        
        return None
    
    def _convert_doc_to_docx(self, doc_path: Path) -> Optional[Path]:
        """
        将 DOC 转换为 DOCX
        
        使用 LibreOffice 进行转换
        """
        output_docx = self.output_dir / f"{doc_path.stem}.docx"
        
        if shutil.which("soffice") or shutil.which("libreoffice"):
            try:
                cmd = [
                    "soffice" if shutil.which("soffice") else "libreoffice",
                    "--headless",
                    "--convert-to", "docx",
                    "--outdir", str(self.output_dir),
                    str(doc_path)
                ]
                
                result = subprocess.run(
                    cmd,
                    timeout=30,
                    capture_output=True,
                    text=True
                )
                
                if result.returncode == 0 and output_docx.exists():
                    logger.info(f"DOC 转 DOCX 成功: {output_docx.name}")
                    return output_docx
                else:
                    logger.warning(f"LibreOffice 转换失败: {result.stderr}")
            except Exception as e:
                logger.warning(f"DOC 转换异常: {e}")
        else:
            logger.error("未找到 LibreOffice，无法转换 DOC 文件")
        
        return None
    
    def _process_docx_direct(self, docx_path: Path) -> DocumentStructure:
        """直接处理 DOCX（备用方案）"""
        doc = Document(str(docx_path))
        
        # 提取文本
        paragraphs = []
        for idx, para in enumerate(doc.paragraphs):
            if para.text.strip():
                paragraphs.append(ParagraphInfo(
                    index=idx,
                    text=para.text.strip()
                ))
        
        # 简单估算页数（每页约 500 字）
        total_chars = sum(len(p.text) for p in paragraphs)
        estimated_pages = max(1, total_chars // 500)
        
        # 构建 Markdown
        md_parts = [p.text for p in paragraphs]
        markdown = "\n\n".join(md_parts)
        
        # 构建页面结构（简化版）
        structure = [PageStructure(
            page=1,
            paragraphs=paragraphs,
            tables=[],
            images=[]
        )]
        
        return DocumentStructure(
            markdown=markdown,
            structure=structure,
            format="docx",
            page_count=estimated_pages
        )

